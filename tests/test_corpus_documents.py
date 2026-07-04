import base64
import json
import zipfile
from datetime import date
from pathlib import Path

import fitz  # type: ignore[import-untyped]
import pytest
import requests
import xlwt  # type: ignore[import-untyped]
from bs4 import BeautifulSoup
from openpyxl import Workbook  # type: ignore[import-untyped]

from axiom_corpus.corpus import documents as documents_module
from axiom_corpus.corpus.artifacts import CorpusArtifactStore
from axiom_corpus.corpus.documents import (
    OFFICIAL_DOCUMENT_BROWSER_USER_AGENT,
    OFFICIAL_DOCUMENT_USER_AGENT,
    OfficialDocumentManifest,
    OfficialDocumentSource,
    _date_text,
    _download_document,
    _download_document_by_curl_ranges,
    _download_document_by_ranges,
    _DownloadedDocument,
    _extract_anchor_range_html_blocks,
    _extract_blocks,
    _extract_doc_blocks,
    _extract_json_html_blocks,
    _extract_json_record_blocks,
    _extract_labeled_html_section_blocks,
    _get_with_retries,
    _infer_source_format,
    _legacy_word_document_text,
    _normalize_text,
    _parse_content_range,
    _parse_curl_header_dump,
    _sanitize_official_document_content,
    _single_scope,
    _validate_citation_path,
    _xlsx_cell_text,
    _xlsx_configured_strings,
    _xlsx_filters,
    _xlsx_row_matches_filters,
    _xlsx_sheet_names,
    extract_official_documents,
    google_drive_download_url,
    official_documents_run_id,
)
from axiom_corpus.corpus.io import load_provisions


def test_google_drive_download_url_converts_file_view():
    url = "https://drive.google.com/file/d/abc123XYZ/view?usp=drive_link"

    assert (
        google_drive_download_url(url) == "https://drive.google.com/uc?export=download&id=abc123XYZ"
    )


def test_normalize_text_removes_invisible_markers():
    assert _normalize_text("\ufeff\u200b Requirements\n\nA rule applies.") == (
        "Requirements\n\nA rule applies."
    )


def test_official_document_manifest_validation_errors(tmp_path: Path) -> None:
    scalar_manifest = tmp_path / "scalar.yaml"
    scalar_manifest.write_text("[]")
    with pytest.raises(ValueError, match="YAML mapping"):
        OfficialDocumentManifest.load(scalar_manifest)

    missing_documents = tmp_path / "missing.yaml"
    missing_documents.write_text("documents: nope")
    with pytest.raises(ValueError, match="documents list"):
        OfficialDocumentManifest.load(missing_documents)

    bad_request = {
        "source_id": "bad",
        "jurisdiction": "be",
        "document_class": "guidance",
        "title": "Bad",
        "source_url": "https://example.test/bad",
        "request": "not-a-mapping",
    }
    with pytest.raises(ValueError, match="request config"):
        OfficialDocumentSource.from_mapping(bad_request)

    bad_extraction = {**bad_request, "request": {}, "extraction": "not-a-mapping"}
    with pytest.raises(ValueError, match="extraction config"):
        OfficialDocumentSource.from_mapping(bad_extraction)

    duplicate = OfficialDocumentManifest(
        documents=(
            OfficialDocumentSource.from_mapping({**bad_request, "request": {}}),
            OfficialDocumentSource.from_mapping({**bad_request, "request": {}}),
        )
    )
    with pytest.raises(ValueError, match="duplicate source_id"):
        duplicate.require_unique_sources()


def test_official_document_scope_requires_single_scope() -> None:
    first = OfficialDocumentSource(
        source_id="a",
        jurisdiction="be",
        document_class="guidance",
        title="A",
        source_url="https://example.test/a",
    )
    second = OfficialDocumentSource(
        source_id="b",
        jurisdiction="be-vlg",
        document_class="guidance",
        title="B",
        source_url="https://example.test/b",
    )

    with pytest.raises(ValueError, match="one jurisdiction"):
        _single_scope((first, second))


def test_run_id_date_text_and_citation_path_helpers() -> None:
    assert official_documents_run_id("2026-07-04", only_source_id="BE Source", limit=2) == (
        "2026-07-04-BE Source-limit-2"
    )
    assert _date_text(None, "fallback") == "fallback"
    assert _date_text(date(2026, 7, 4), "fallback") == "2026-07-04"
    assert _date_text("2026-Q1", "fallback") == "2026-Q1"
    assert (
        _validate_citation_path(
            "be/guidance/onss/dmfa",
            jurisdiction="be",
            document_class="guidance",
        )
        == "be/guidance/onss/dmfa"
    )
    with pytest.raises(ValueError, match="citation_path must start"):
        _validate_citation_path(
            "be/statute/onss",
            jurisdiction="be",
            document_class="guidance",
        )


def test_infer_source_format_handles_common_official_downloads(tmp_path: Path) -> None:
    def source(source_url: str, *, source_format: str | None = None) -> OfficialDocumentSource:
        return OfficialDocumentSource(
            source_id="doc",
            jurisdiction="be",
            document_class="guidance",
            title="Document",
            source_url=source_url,
            source_format=source_format,
        )

    def downloaded(
        official_source: OfficialDocumentSource,
        content: bytes,
        content_type: str | None,
    ) -> _DownloadedDocument:
        return _DownloadedDocument(
            source=official_source,
            content=content,
            content_type=content_type,
            final_url=official_source.source_url,
        )

    explicit = source("https://example.test/file.bin", source_format="HTML")
    assert _infer_source_format(explicit, downloaded(explicit, b"ignored", None)) == "html"

    pdf = source("https://example.test/file")
    assert _infer_source_format(pdf, downloaded(pdf, b"%PDF-1.7", None)) == "pdf"
    html = source("https://example.test/file")
    assert _infer_source_format(html, downloaded(html, b"<!doctype html>", None)) == "html"
    xls = source("https://example.test/file.xls?download=1")
    assert _infer_source_format(xls, downloaded(xls, b"legacy", None)) == "xls"
    doc = source("https://example.test/file")
    assert (
        _infer_source_format(
            doc,
            downloaded(doc, b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1legacy", None),
        )
        == "doc"
    )

    docx_path = tmp_path / "word.zip"
    with zipfile.ZipFile(docx_path, "w") as archive:
        archive.writestr("word/document.xml", "<w:document/>")
    docx = source("https://example.test/file")
    assert _infer_source_format(docx, downloaded(docx, docx_path.read_bytes(), None)) == "docx"

    xlsx_path = tmp_path / "sheet.zip"
    with zipfile.ZipFile(xlsx_path, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
    xlsx = source("https://example.test/file")
    assert _infer_source_format(xlsx, downloaded(xlsx, xlsx_path.read_bytes(), None)) == "xlsx"

    unknown = source("https://example.test/file")
    with pytest.raises(ValueError, match="cannot infer"):
        _infer_source_format(unknown, downloaded(unknown, b"plain text", "text/plain"))


def test_sanitize_official_document_content_redacts_public_tokens() -> None:
    token = b"pk.abc_123.DEF-456"

    sanitized_html = _sanitize_official_document_content(b"<script>" + token + b"</script>", "html")
    sanitized_json = _sanitize_official_document_content(b'{"token":"' + token + b'"}', "json")

    assert token not in sanitized_html
    assert b"[redacted-mapbox-public-token]" in sanitized_html
    assert token not in sanitized_json
    assert _sanitize_official_document_content(token, "pdf") == token


def test_extract_blocks_rejects_unsupported_and_bad_json_config() -> None:
    with pytest.raises(ValueError, match="unsupported official document source_format"):
        _extract_blocks(
            b"body",
            "csv",
            source_url="https://example.test/file.csv",
            title="File",
            extraction=None,
        )

    with pytest.raises(ValueError, match="requires json_html_field"):
        _extract_blocks(
            b'{"html": "<p>body</p>"}',
            "json",
            source_url="https://example.test/file.json",
            title="File",
            extraction={},
        )


def test_extract_doc_blocks_rejects_segmentation_and_empty_doc(monkeypatch) -> None:
    with pytest.raises(ValueError, match="only supports single_block"):
        _extract_doc_blocks(
            b"content",
            title="Legacy",
            extraction={"segmentation": "labeled_sections"},
        )

    monkeypatch.setattr(documents_module, "_legacy_word_document_text", lambda content: "")
    assert _extract_doc_blocks(b"content", title="Legacy", extraction=None) == ()


def test_xlsx_helper_configuration_filters_and_cells() -> None:
    assert _xlsx_sheet_names(["A", "B"], {}) == ("A", "B")
    assert _xlsx_sheet_names(["A"], {"sheet": "A"}) == ("A",)
    with pytest.raises(ValueError, match="sheet not found"):
        _xlsx_sheet_names(["A"], {"sheets": ["B"]})

    assert _xlsx_configured_strings(None) == ()
    assert _xlsx_configured_strings("A") == ("A",)
    assert _xlsx_configured_strings(["A", 2]) == ("A", "2")
    with pytest.raises(ValueError, match="string or list"):
        _xlsx_configured_strings({"bad": "value"})

    assert _xlsx_filters({"Year": [2024, 2025], "Active": True}) == {
        "Year": ("2024", "2025"),
        "Active": ("TRUE",),
    }
    with pytest.raises(ValueError, match="filters must be a mapping"):
        _xlsx_filters(["bad"])

    index = {"Year": 0, "Amount": 1}
    assert _xlsx_row_matches_filters(("2024", "10"), index=index, filters={"Year": ("2024",)})
    assert not _xlsx_row_matches_filters(
        ("2025", "10"),
        index=index,
        filters={"Year": ("2024",)},
    )
    assert not _xlsx_row_matches_filters(("2024",), index=index, filters={"Amount": ("10",)})
    with pytest.raises(ValueError, match="filter column not found"):
        _xlsx_row_matches_filters(("2024",), index=index, filters={"Missing": ("x",)})

    assert _xlsx_cell_text(None) == ""
    assert _xlsx_cell_text(True) == "TRUE"
    assert _xlsx_cell_text(date(2026, 7, 4)) == "2026-07-04"
    assert _xlsx_cell_text(4) == "4"
    assert _xlsx_cell_text(4.0) == "4"
    assert _xlsx_cell_text(4.25) == "4.25"
    assert _xlsx_cell_text(" A\n\nB ") == "A\n\nB"


def test_extract_json_record_blocks_handles_nested_fields_and_filters() -> None:
    content = json.dumps(
        {
            "items": [
                "not a record",
                {
                    "status": "draft",
                    "text": {"body": "Draft text"},
                    "label": "D",
                },
                {
                    "id": 12,
                    "status": "current",
                    "text": {"body": "Line 1\n\nLine 2"},
                    "label": "A:1",
                    "heading": "",
                    "kind": {"name": "Rule"},
                },
                {
                    "id": 13,
                    "status": "current",
                    "text": {"body": "   "},
                    "label": "A:2",
                },
            ]
        }
    ).encode("utf-8")

    blocks = _extract_json_record_blocks(
        content,
        source_url="https://example.test/api",
        fallback_title="Fallback Title",
        extraction={
            "json_records_path": "items",
            "json_record_text_field": "text.body",
            "json_record_text_is_html": False,
            "json_record_label_field": "label",
            "json_record_heading_field": "heading",
            "json_record_kind_field": "kind.name",
            "json_record_status_field": "status",
            "json_record_include_statuses": "current",
            "json_record_metadata_fields": "id",
        },
    )

    assert len(blocks) == 1
    block = blocks[0]
    assert block.kind == "rule"
    assert block.heading == "A:1 Fallback Title"
    assert block.body == "Line 1\n\nLine 2"
    assert block.metadata["id"] == 12
    assert block.metadata["citation_suffix"] == "A:1"

    with pytest.raises(ValueError, match="requires json_record_text_field"):
        _extract_json_record_blocks(
            b"[]",
            source_url="https://example.test/api",
            fallback_title=None,
            extraction={},
        )
    with pytest.raises(ValueError, match="must resolve to a list"):
        _extract_json_record_blocks(
            b'{"items": {"not": "a-list"}}',
            source_url="https://example.test/api",
            fallback_title=None,
            extraction={
                "json_records_path": "items",
                "json_record_text_field": "text",
            },
        )


def test_extract_json_record_blocks_rejects_bad_field_configs() -> None:
    content = b"[]"
    for key in (
        "json_record_label_field",
        "json_record_heading_field",
        "json_record_kind_field",
        "json_record_status_field",
    ):
        with pytest.raises(ValueError, match="must be a string"):
            _extract_json_record_blocks(
                content,
                source_url="https://example.test/api",
                fallback_title=None,
                extraction={
                    "json_record_text_field": "text",
                    key: 1,
                },
            )

    blocks = _extract_json_record_blocks(
        json.dumps(
            [
                {
                    "status": "skip",
                    "text": "<p>Skip</p>",
                },
                {
                    "status": "keep",
                    "text": "<script>ignored</script>",
                },
                {
                    "status": "keep",
                    "text": "<p>Body</p>",
                },
            ]
        ).encode("utf-8"),
        source_url="https://example.test/api",
        fallback_title=None,
        extraction={
            "json_record_text_field": "text",
            "json_record_status_field": "status",
            "json_record_exclude_statuses": "skip",
        },
    )
    assert len(blocks) == 1
    assert blocks[0].heading == "Record"
    assert blocks[0].body == "Body"


def test_extract_json_html_blocks_errors_and_empty_single_block() -> None:
    with pytest.raises(ValueError, match="did not resolve to HTML text"):
        _extract_json_html_blocks(
            b'{"html": "   "}',
            source_url="https://example.test/api",
            fallback_title="JSON",
            extraction={"json_html_field": "html"},
        )
    with pytest.raises(ValueError, match="json path did not resolve"):
        _extract_json_html_blocks(
            b'{"output": {}}',
            source_url="https://example.test/api",
            fallback_title="JSON",
            extraction={"json_html_field": "output.html"},
        )
    assert (
        _extract_json_html_blocks(
            b'{"html": "<script>ignored</script>"}',
            source_url="https://example.test/api",
            fallback_title="JSON",
            extraction={"json_html_field": "html", "json_html_as_single_block": True},
        )
        == ()
    )


def test_extract_anchor_range_html_blocks_validates_config_and_extracts_ranges() -> None:
    soup = BeautifulSoup(
        """
        <main>
          <a id="a">Article A</a>
          <p>First body.</p>
          <!-- ignored -->
          <a id="b">Article B</a>
          <p>Second body.</p>
          <a id="c">Article C</a>
        </main>
        """,
        "html.parser",
    )
    root = soup.main
    assert root is not None

    blocks = _extract_anchor_range_html_blocks(
        root,
        title="Document",
        source_url="https://example.test/html",
        extraction={
            "anchor_ranges": [
                {
                    "html_start_selector": "#a",
                    "html_stop_selector": "#b",
                    "section_label": "article-a",
                },
                {
                    "html_start_selector": "#b",
                    "html_stop_selector": "#c",
                    "section_heading": "Article B",
                },
            ]
        },
    )

    assert [block.ordinal for block in blocks] == [1, 2]
    assert blocks[0].metadata["citation_suffix"] == "article-a"
    assert "First body" in blocks[0].body
    assert "Second body" not in blocks[0].body
    assert blocks[1].heading == "Article B"
    assert "Second body" in blocks[1].body

    with pytest.raises(ValueError, match="anchor_ranges must be a list"):
        _extract_anchor_range_html_blocks(
            root,
            title=None,
            source_url="https://example.test/html",
            extraction={"anchor_ranges": "bad"},
        )
    with pytest.raises(ValueError, match="entries must be mappings"):
        _extract_anchor_range_html_blocks(
            root,
            title=None,
            source_url="https://example.test/html",
            extraction={"anchor_ranges": ["bad"]},
        )
    with pytest.raises(ValueError, match="requires html_start_selector"):
        _extract_anchor_range_html_blocks(
            root,
            title=None,
            source_url="https://example.test/html",
            extraction={},
        )
    with pytest.raises(ValueError, match="start selector did not match"):
        _extract_anchor_range_html_blocks(
            root,
            title=None,
            source_url="https://example.test/html",
            extraction={"html_start_selector": "#missing"},
        )
    with pytest.raises(ValueError, match="non-empty string"):
        _extract_anchor_range_html_blocks(
            root,
            title=None,
            source_url="https://example.test/html",
            extraction={"html_start_selector": "#a", "html_stop_selector": 1},
        )
    with pytest.raises(ValueError, match="stop selector did not match"):
        _extract_anchor_range_html_blocks(
            root,
            title=None,
            source_url="https://example.test/html",
            extraction={"html_start_selector": "#a", "html_stop_selector": "#missing"},
        )


def test_extract_labeled_html_sections_label_only_and_validation() -> None:
    soup = BeautifulSoup(
        """
        <main>
          <p> </p>
          <p>1</p>
          <p>Body text.</p>
          <p>STOP</p>
          <p>Ignored text.</p>
        </main>
        """,
        "html.parser",
    )
    root = soup.main
    assert root is not None

    with pytest.raises(ValueError, match="requires section_heading_pattern"):
        _extract_labeled_html_section_blocks(
            root,
            title="Document",
            source_url="https://example.test/html",
            extraction={},
        )
    with pytest.raises(ValueError, match="section_label_replacements must be a mapping"):
        _extract_labeled_html_section_blocks(
            root,
            title="Document",
            source_url="https://example.test/html",
            extraction={
                "section_label_pattern": r"^(?P<label>\d+)$",
                "section_label_replacements": ["bad"],
            },
        )

    blocks = _extract_labeled_html_section_blocks(
        root,
        title="Document",
        source_url="https://example.test/html",
        extraction={
            "section_label_pattern": r"^(?P<label>\d+)$",
            "section_label_replacements": {"1": "article-one"},
            "stop_text_pattern": "^STOP$",
        },
    )

    assert len(blocks) == 1
    assert blocks[0].heading == "article-one"
    assert blocks[0].body == "Body text."
    assert blocks[0].metadata["citation_suffix"] == "article-one"


def test_download_document_retries_browser_user_agent_on_forbidden():
    class FakeResponse:
        def __init__(self, status_code: int, content: bytes = b""):
            self.status_code = status_code
            self.content = content
            self.headers = {"content-type": "application/pdf"}
            self.url = "https://example.test/doc.pdf"

        def close(self):
            return None

        def raise_for_status(self):
            if self.status_code >= 400:
                raise AssertionError("final response should not be an error")

    class FakeSession:
        def __init__(self):
            self.headers = {"User-Agent": OFFICIAL_DOCUMENT_USER_AGENT}
            self.calls: list[dict[str, str]] = []

        def get(self, url, *, headers=None, timeout=None, allow_redirects=None, verify=None):
            del url, timeout, allow_redirects, verify
            self.calls.append(dict(headers or self.headers))
            if len(self.calls) == 1:
                return FakeResponse(403)
            return FakeResponse(200, b"%PDF-1.7")

    source = OfficialDocumentSource(
        source_id="doc",
        jurisdiction="us-test",
        document_class="form",
        title="Document",
        source_url="https://example.test/doc.pdf",
    )
    session = FakeSession()

    downloaded = _download_document(source, session=session)  # pyright: ignore[reportPrivateUsage]

    assert downloaded.content == b"%PDF-1.7"
    assert session.calls[0]["User-Agent"] == OFFICIAL_DOCUMENT_USER_AGENT
    assert session.calls[1]["User-Agent"] == OFFICIAL_DOCUMENT_BROWSER_USER_AGENT


def test_download_document_retries_browser_user_agent_on_declared_pdf_html_challenge():
    class FakeResponse:
        def __init__(self, content: bytes, content_type: str):
            self.status_code = 200
            self.content = content
            self.headers = {"content-type": content_type}
            self.url = "https://example.test/doc.pdf"

        def close(self):
            return None

        def raise_for_status(self):
            return None

        def iter_content(self, chunk_size: int):
            del chunk_size
            yield self.content

    class FakeSession:
        def __init__(self):
            self.headers = {"User-Agent": OFFICIAL_DOCUMENT_USER_AGENT}
            self.calls: list[dict[str, str]] = []

        def get(self, url, *, headers=None, timeout=None, allow_redirects=None, verify=None):
            del url, timeout, allow_redirects, verify
            self.calls.append(dict(headers or self.headers))
            if len(self.calls) == 1:
                return FakeResponse(b"<html><body>challenge</body></html>", "text/html")
            return FakeResponse(b"%PDF-1.7", "application/pdf")

    source = OfficialDocumentSource(
        source_id="doc",
        jurisdiction="us-test",
        document_class="form",
        title="Document",
        source_url="https://example.test/doc.pdf",
        source_format="pdf",
    )
    session = FakeSession()

    downloaded = _download_document(source, session=session)  # pyright: ignore[reportPrivateUsage]

    assert downloaded.content == b"%PDF-1.7"
    assert session.calls[0]["User-Agent"] == OFFICIAL_DOCUMENT_USER_AGENT
    assert session.calls[1]["User-Agent"] == OFFICIAL_DOCUMENT_BROWSER_USER_AGENT


def test_download_document_retries_transient_request_errors(monkeypatch):
    class FakeResponse:
        status_code = 200
        content = b"<html><body>ok</body></html>"
        headers = {"content-type": "text/html"}
        url = "https://example.test/doc.html"

        def close(self):
            return None

        def raise_for_status(self):
            return None

        def iter_content(self, chunk_size: int):
            del chunk_size
            yield self.content

    class FakeSession:
        def __init__(self):
            self.headers = {"User-Agent": OFFICIAL_DOCUMENT_USER_AGENT}
            self.calls = 0

        def get(self, url, *, headers=None, timeout=None, allow_redirects=None, verify=None):
            del url, headers, timeout, allow_redirects, verify
            self.calls += 1
            if self.calls == 1:
                raise requests.exceptions.ChunkedEncodingError("connection reset")
            return FakeResponse()

    monkeypatch.setattr(documents_module.time, "sleep", lambda seconds: None)
    source = OfficialDocumentSource(
        source_id="doc",
        jurisdiction="us-test",
        document_class="manual",
        title="Document",
        source_url="https://example.test/doc.html",
    )
    session = FakeSession()

    downloaded = _download_document(source, session=session)  # pyright: ignore[reportPrivateUsage]

    assert downloaded.content == b"<html><body>ok</body></html>"
    assert session.calls == 2


def test_download_document_can_disable_tls_verification():
    class FakeResponse:
        status_code = 200
        content = b'{"ok": true}'
        headers = {"content-type": "application/json"}
        url = "https://example.test/doc.json"

        def close(self):
            return None

        def raise_for_status(self):
            return None

        def iter_content(self, chunk_size: int):
            del chunk_size
            yield self.content

    class FakeSession:
        def __init__(self):
            self.headers = {"User-Agent": OFFICIAL_DOCUMENT_USER_AGENT}
            self.verify_values: list[bool | None] = []

        def get(self, url, *, headers=None, timeout=None, allow_redirects=None, verify=None):
            del url, headers, timeout, allow_redirects
            self.verify_values.append(verify)
            return FakeResponse()

    source = OfficialDocumentSource(
        source_id="doc",
        jurisdiction="us-test",
        document_class="regulation",
        title="Document",
        source_url="https://example.test/doc.json",
        request={"verify_tls": False},
    )
    session = FakeSession()

    downloaded = _download_document(source, session=session)  # pyright: ignore[reportPrivateUsage]

    assert downloaded.content == b'{"ok": true}'
    assert session.verify_values == [False]


def test_download_document_supports_range_fetch():
    payload = b"%PDF-1.7\nexample"

    class FakeResponse:
        def __init__(
            self,
            content: bytes,
            *,
            start: int,
            end: int,
            total: int,
        ):
            self.status_code = 206
            self.content = content
            self.headers = {
                "content-type": "application/pdf",
                "content-range": f"bytes {start}-{end}/{total}",
            }
            self.url = "https://example.test/doc.pdf"

        def close(self):
            return None

        def raise_for_status(self):
            return None

        def iter_content(self, chunk_size: int):
            del chunk_size
            yield self.content

    class FakeSession:
        def __init__(self):
            self.headers = {"User-Agent": OFFICIAL_DOCUMENT_USER_AGENT}
            self.ranges: list[str] = []
            self.user_agents: list[str] = []

        def get(
            self,
            url,
            *,
            headers=None,
            timeout=None,
            allow_redirects=None,
            verify=None,
            stream=None,
        ):
            del url, timeout, allow_redirects, verify
            assert stream is True
            request_headers = headers or {}
            range_header = str(request_headers["Range"])
            self.ranges.append(range_header)
            self.user_agents.append(str(request_headers["User-Agent"]))
            requested = range_header.removeprefix("bytes=")
            start_text, end_text = requested.split("-", 1)
            start = int(start_text)
            end = min(int(end_text), len(payload) - 1)
            return FakeResponse(
                payload[start : end + 1],
                start=start,
                end=end,
                total=len(payload),
            )

    source = OfficialDocumentSource(
        source_id="doc",
        jurisdiction="us-test",
        document_class="form",
        title="Document",
        source_url="https://example.test/doc.pdf",
        source_format="pdf",
        request={"range_fetch": True, "range_chunk_size": 5, "browser_user_agent": True},
    )
    session = FakeSession()

    downloaded = _download_document(source, session=session)  # pyright: ignore[reportPrivateUsage]

    assert downloaded.content == payload
    assert session.ranges == ["bytes=0-4", "bytes=5-9", "bytes=10-14", "bytes=15-19"]
    assert session.user_agents == [OFFICIAL_DOCUMENT_BROWSER_USER_AGENT] * 4


def test_download_document_supports_curl_range_backend(monkeypatch):
    payload = b"%PDF-1.7\nexample"
    commands: list[list[str]] = []

    def fake_run(command, check):
        assert check is True
        commands.append(list(command))
        range_index = command.index("-H") + 1
        while not str(command[range_index]).startswith("Range: "):
            range_index = command.index("-H", range_index + 1) + 1
        requested = str(command[range_index]).removeprefix("Range: bytes=")
        start_text, end_text = requested.split("-", 1)
        start = int(start_text)
        end = min(int(end_text), len(payload) - 1)
        header_path = Path(command[command.index("--dump-header") + 1])
        body_path = Path(command[command.index("--output") + 1])
        header_path.write_text(
            "\r\n".join(
                [
                    "HTTP/1.1 206 Partial Content",
                    "Content-Type: application/pdf",
                    f"Content-Range: bytes {start}-{end}/{len(payload)}",
                    "",
                    "",
                ]
            )
        )
        body_path.write_bytes(payload[start : end + 1])

    monkeypatch.setattr(documents_module.subprocess, "run", fake_run)
    source = OfficialDocumentSource(
        source_id="doc",
        jurisdiction="us-test",
        document_class="form",
        title="Document",
        source_url="https://example.test/doc.pdf",
        source_format="pdf",
        request={
            "range_fetch": True,
            "range_backend": "curl",
            "range_chunk_size": 5,
            "browser_user_agent": True,
            "verify_tls": False,
        },
    )

    downloaded = _download_document(source, session=requests.Session())  # pyright: ignore[reportPrivateUsage]

    assert downloaded.content == payload
    assert [command[command.index("-A") + 1] for command in commands] == [
        OFFICIAL_DOCUMENT_BROWSER_USER_AGENT
    ] * 4
    assert all("--insecure" in command for command in commands)


class FakeRangeResponse:
    def __init__(
        self,
        status_code: int,
        content: bytes,
        *,
        content_range: str | None = None,
    ) -> None:
        self.status_code = status_code
        self.content = content
        self.url = "https://example.test/file"
        self.headers = {"content-type": "application/octet-stream"}
        if content_range is not None:
            self.headers["content-range"] = content_range
        self.closed = False

    def raise_for_status(self):
        return None

    def iter_content(self, chunk_size: int):
        del chunk_size
        yield self.content

    def close(self):
        self.closed = True


class FakeRangeSession:
    def __init__(self, responses: list[FakeRangeResponse]) -> None:
        self.responses = responses
        self.requests: list[dict[str, str]] = []

    def get(self, url, **kwargs):
        del url
        self.requests.append(dict(kwargs.get("headers") or {}))
        return self.responses.pop(0)


def test_get_with_retries_reraises_final_request_exception(monkeypatch) -> None:
    class FailingSession:
        def get(self, *args, **kwargs):
            del args, kwargs
            raise requests.RequestException("offline")

    monkeypatch.setattr(documents_module.time, "sleep", lambda seconds: None)
    with pytest.raises(requests.RequestException, match="offline"):
        _get_with_retries(FailingSession(), "https://example.test/file")


def test_download_document_by_ranges_handles_full_response() -> None:
    source = OfficialDocumentSource(
        source_id="range-full",
        jurisdiction="be",
        document_class="guidance",
        title="Range",
        source_url="https://example.test/file",
    )
    session = FakeRangeSession([FakeRangeResponse(200, b"complete")])

    downloaded = _download_document_by_ranges(
        source,
        "https://example.test/file",
        session=session,
        headers={"Accept": "*/*"},
        verify=True,
        chunk_size=4,
    )

    assert downloaded.content == b"complete"
    assert session.requests == [{"Accept": "*/*", "Range": "bytes=0-3"}]


def test_download_document_by_ranges_rejects_bad_status_start_and_size() -> None:
    source = OfficialDocumentSource(
        source_id="range-error",
        jurisdiction="be",
        document_class="guidance",
        title="Range",
        source_url="https://example.test/file",
    )

    with pytest.raises(RuntimeError, match="HTTP 416"):
        _download_document_by_ranges(
            source,
            "https://example.test/file",
            session=FakeRangeSession([FakeRangeResponse(416, b"")]),
            headers=None,
            verify=True,
            chunk_size=4,
        )

    with pytest.raises(RuntimeError, match="unexpected start 1"):
        _download_document_by_ranges(
            source,
            "https://example.test/file",
            session=FakeRangeSession(
                [FakeRangeResponse(206, b"abc", content_range="bytes 1-3/10")]
            ),
            headers=None,
            verify=True,
            chunk_size=4,
        )

    with pytest.raises(ValueError, match="positive"):
        _download_document_by_ranges(
            source,
            "https://example.test/file",
            session=FakeRangeSession([]),
            headers=None,
            verify=True,
            chunk_size=0,
        )


def test_parse_content_range_and_curl_headers_errors() -> None:
    assert _parse_content_range("bytes 0-3/*") == (0, 3, None)
    with pytest.raises(RuntimeError, match="invalid Content-Range"):
        _parse_content_range("not-a-range")
    with pytest.raises(RuntimeError, match="did not include"):
        _parse_curl_header_dump("")
    with pytest.raises(RuntimeError, match="invalid curl HTTP status"):
        _parse_curl_header_dump("HTTP/1.1\ncontent-type: text/plain")


def test_download_document_by_ranges_rejects_full_after_partial_and_empty_unknown_total() -> None:
    source = OfficialDocumentSource(
        source_id="range-partial",
        jurisdiction="be",
        document_class="guidance",
        title="Range",
        source_url="https://example.test/file",
    )

    with pytest.raises(RuntimeError, match="full response after partial chunks"):
        _download_document_by_ranges(
            source,
            "https://example.test/file",
            session=FakeRangeSession(
                [
                    FakeRangeResponse(206, b"ab", content_range="bytes 0-1/*"),
                    FakeRangeResponse(200, b"complete"),
                ]
            ),
            headers=None,
            verify=True,
            chunk_size=2,
        )

    downloaded = _download_document_by_ranges(
        source,
        "https://example.test/file",
        session=FakeRangeSession([FakeRangeResponse(206, b"", content_range="bytes 0-0/*")]),
        headers=None,
        verify=True,
        chunk_size=2,
    )
    assert downloaded.content == b""


def test_get_with_retries_retries_retryable_status(monkeypatch) -> None:
    class RetryableResponse:
        def __init__(self, status_code: int) -> None:
            self.status_code = status_code
            self.closed = False

        def close(self) -> None:
            self.closed = True

    class RetryableSession:
        def __init__(self) -> None:
            self.responses = [RetryableResponse(503), RetryableResponse(200)]
            self.calls: list[dict[str, object]] = []

        def get(self, url, **kwargs):
            del url
            self.calls.append(dict(kwargs))
            return self.responses.pop(0)

    monkeypatch.setattr(documents_module.time, "sleep", lambda seconds: None)
    session = RetryableSession()

    response = _get_with_retries(
        session,
        "https://example.test/file",
        headers={"Accept": "text/plain"},
        verify=False,
        stream=True,
    )

    assert response.status_code == 200
    assert session.calls[0]["stream"] is True
    assert session.calls[0]["verify"] is False


def test_download_document_by_curl_ranges_handles_error_edges(monkeypatch) -> None:
    source = OfficialDocumentSource(
        source_id="curl-range",
        jurisdiction="be",
        document_class="guidance",
        title="Curl Range",
        source_url="https://example.test/file",
    )
    responses: list[tuple[int, str | None, bytes]] = []
    commands: list[list[str]] = []

    def fake_run(command, check):
        assert check is True
        commands.append(list(command))
        status, content_range, body = responses.pop(0)
        header_path = Path(command[command.index("--dump-header") + 1])
        body_path = Path(command[command.index("--output") + 1])
        header_lines = [
            f"HTTP/1.1 {status} Status",
            "Content-Type: application/octet-stream",
            "Header without colon",
        ]
        if content_range is not None:
            header_lines.append(f"Content-Range: {content_range}")
        header_path.write_text("\r\n".join([*header_lines, "", ""]))
        body_path.write_bytes(body)

    monkeypatch.setattr(documents_module.subprocess, "run", fake_run)

    with pytest.raises(ValueError, match="positive"):
        _download_document_by_curl_ranges(
            source,
            "https://example.test/file",
            headers=None,
            verify=True,
            chunk_size=0,
        )

    responses[:] = [(500, None, b"error")]
    with pytest.raises(RuntimeError, match="HTTP 500"):
        _download_document_by_curl_ranges(
            source,
            "https://example.test/file",
            headers={"User-Agent": "Custom UA", "Accept": "text/plain"},
            verify=True,
            chunk_size=2,
        )
    assert "Accept: text/plain" in commands[-1]

    responses[:] = [(206, "bytes 1-2/4", b"bc")]
    with pytest.raises(RuntimeError, match="unexpected start 1"):
        _download_document_by_curl_ranges(
            source,
            "https://example.test/file",
            headers=None,
            verify=True,
            chunk_size=2,
        )

    responses[:] = [
        (206, "bytes 0-1/4", b"ab"),
        (200, None, b"complete"),
    ]
    with pytest.raises(RuntimeError, match="full response after partial chunks"):
        _download_document_by_curl_ranges(
            source,
            "https://example.test/file",
            headers=None,
            verify=True,
            chunk_size=2,
        )

    responses[:] = [(206, "bytes 0-0/*", b"")]
    downloaded = _download_document_by_curl_ranges(
        source,
        "https://example.test/file",
        headers=None,
        verify=True,
        chunk_size=2,
    )
    assert downloaded.content == b""


def test_legacy_word_document_text_uses_tool_and_reports_errors(monkeypatch) -> None:
    with pytest.raises(RuntimeError, match="requires textutil"):
        monkeypatch.setattr(documents_module.shutil, "which", lambda name: None)
        _legacy_word_document_text(b"doc")

    class Result:
        def __init__(self, returncode: int, stdout: str, stderr: str = "") -> None:
            self.returncode = returncode
            self.stdout = stdout
            self.stderr = stderr

    monkeypatch.setattr(
        documents_module.shutil,
        "which",
        lambda name: "/usr/bin/textutil" if name == "textutil" else None,
    )
    next_result = Result(1, "", "converter failed")

    def fake_run(command, check, capture_output, encoding, errors):
        assert command[0] == "textutil"
        assert check is False
        assert capture_output is True
        assert encoding == "utf-8"
        assert errors == "replace"
        return next_result

    monkeypatch.setattr(documents_module.subprocess, "run", fake_run)
    with pytest.raises(RuntimeError, match="converter failed"):
        _legacy_word_document_text(b"doc")

    next_result = Result(0, " A\n\nB ")
    assert _legacy_word_document_text(b"doc") == "A\n\nB"


def test_extract_official_documents_from_local_html_and_pdf(tmp_path):
    html_path = tmp_path / "snap.html"
    long_html_text = "Long eligibility detail. " * 180
    html_path.write_text(
        f"""
        <html>
          <head><title>Ignored browser title</title></head>
          <body>
            <nav>Navigation should not become its own block.</nav>
            <main>
              <h1>Colorado SNAP Policy</h1>
              <h2>Eligibility</h2>
              <p>Households may qualify based on income and household size.</p>
              <p>{long_html_text}</p>
              <ul><li>County departments determine eligibility.</li></ul>
            </main>
          </body>
        </html>
        """
    )
    pdf_path = tmp_path / "waiver.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "SNAP waiver approval\nApproved for a limited area.")
    document.save(pdf_path)
    document.close()
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: co-snap-page
    jurisdiction: us-co
    document_class: policy
    title: Colorado SNAP Policy
    source_url: https://cdhs.colorado.gov/snap
    citation_path: us-co/policy/cdhs/snap
    source_format: html
    local_path: {json.dumps(str(html_path))}
    metadata:
      source_authority: Colorado Department of Human Services
      document_subtype: agency_page
  - source_id: co-snap-waiver
    jurisdiction: us-co
    document_class: policy
    title: Colorado SNAP Waiver Approval
    source_url: https://www.fns.usda.gov/example.pdf
    source_format: pdf
    local_path: {json.dumps(str(pdf_path))}
    metadata:
      source_authority: USDA Food and Nutrition Service
      document_subtype: waiver_approval
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-04-30",
        source_as_of="2026-04-30",
    )

    assert report.document_count == 2
    assert report.block_count == 2
    assert report.provisions_written == 4
    assert report.coverage.complete
    assert report.inventory_path.exists()
    assert len(report.source_paths) == 2

    inventory = json.loads(report.inventory_path.read_text())
    assert [item["citation_path"] for item in inventory["items"]] == [
        "us-co/policy/cdhs/snap",
        "us-co/policy/cdhs/snap/block-1",
        "us-co/policy/co-snap-waiver",
        "us-co/policy/co-snap-waiver/page-1",
    ]
    records = load_provisions(report.provisions_path)
    page_record = next(record for record in records if record.kind == "page")
    assert page_record.body is not None
    assert "Approved for a limited area" in page_record.body
    assert page_record.source_id == "co-snap-waiver"
    assert page_record.source_document_id is None
    assert page_record.metadata is not None
    assert page_record.metadata["document_subtype"] == "waiver_approval"


def test_extract_official_documents_from_filtered_xlsx_rows(tmp_path: Path) -> None:
    workbook_path = tmp_path / "index.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TA_ALL_BASE_YR"
    sheet.append(["NM_YR", "NM_MTH", "MS_HLTH_IDX", "NM_BASE_YR"])
    sheet.append([2020, 4, 110.22, 2013])
    sheet.append([2024, 4, 130.85, 2013])
    sheet.append([2024, 5, 131.12, 2013])
    workbook.save(workbook_path)
    workbook.close()
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: statbel-health-index
    jurisdiction: be
    document_class: guidance
    title: Statbel health index workbook
    source_url: https://statbel.fgov.be/en/open-data/consumer-price-index-and-health-index
    citation_path: be/guidance/statbel/health-index
    source_format: xlsx
    local_path: {json.dumps(str(workbook_path))}
    extraction:
      xlsx_sheet: TA_ALL_BASE_YR
      xlsx_columns:
        - NM_YR
        - NM_MTH
        - MS_HLTH_IDX
        - NM_BASE_YR
      xlsx_filters:
        NM_BASE_YR: 2013
        NM_MTH: 4
        NM_YR:
          - 2020
          - 2024
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-07-03-statbel-health-index",
    )

    assert report.block_count == 1
    records = load_provisions(report.provisions_path)
    assert records[1].citation_path == "be/guidance/statbel/health-index/sheet-1"
    assert records[1].body is not None
    assert "2020 | 4 | 110.22 | 2013" in records[1].body
    assert "2024 | 4 | 130.85 | 2013" in records[1].body
    assert "131.12" not in records[1].body


def test_extract_official_documents_from_legacy_xls_rows(tmp_path: Path) -> None:
    workbook_path = tmp_path / "ffe.xls"
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("FFE 2025")
    for column, header in enumerate(["Category", "Employers", "Worker count", "Workers", "Rate"]):
        sheet.write(0, column, header)
    sheet.write(1, 0, "Cotisations de base FFE 2025")
    sheet.write(1, 1, "Employeurs avec finalites industrielles ou commerciales")
    sheet.write(1, 2, "a) en moyenne moins de 20 travailleurs")
    sheet.write(1, 3, "tout le personnel")
    sheet.write(1, 4, "0,17% (0,18%)")
    workbook.save(str(workbook_path))

    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: be-onss-ffe-rates
    jurisdiction: be
    document_class: guidance
    title: ONSS FFE rates
    source_url: https://www.socialsecurity.be/site_fr/Infos/instructs/documents/xls/FFE_2025.xls
    citation_path: be/guidance/onss/dmfa/2025-q1/company-closing-fund-rates
    source_format: xls
    local_path: {json.dumps(str(workbook_path))}
    extraction:
      xls_sheet: FFE 2025
      xls_columns:
        - Category
        - Worker count
        - Rate
      citation_suffix: ffe-2025-rates
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-07-04-be-onss-ffe",
    )

    assert report.block_count == 1
    records = load_provisions(report.provisions_path)
    assert records[1].citation_path == (
        "be/guidance/onss/dmfa/2025-q1/company-closing-fund-rates/ffe-2025-rates"
    )
    assert records[1].body is not None
    assert "a) en moyenne moins de 20 travailleurs | 0,17% (0,18%)" in records[1].body


def test_extract_official_documents_can_ocr_scanned_pdf_sections(tmp_path, monkeypatch):
    pdf_path = tmp_path / "scanned-rule.pdf"
    document = fitz.open()
    document.new_page()
    document.save(pdf_path)
    document.close()
    ocr_calls = 0

    def fake_ocr_page(page, *, extraction):
        nonlocal ocr_calls
        del page
        ocr_calls += 1
        assert extraction["ocr"] is True
        return "\n".join(
            [
                "49-001",
                "FIRST SECTION",
                "OCR body text.",
                "49-002 SECOND SECTION",
                "More OCR body text.",
            ]
        )

    monkeypatch.setattr(documents_module, "_ocr_pdf_page_text", fake_ocr_page)
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: scanned-rule
    jurisdiction: us-test
    document_class: regulation
    title: Scanned Rule
    source_url: https://example.test/scanned-rule.pdf
    citation_path: us-test/regulation/scanned-rule
    source_format: pdf
    local_path: {json.dumps(str(pdf_path))}
    extraction:
      ocr: true
      segmentation: labeled_sections
      section_heading_pattern: '^(?P<label>49-[0-9]{{3}})\\s+(?P<heading>[A-Z ]+)$'
      section_label_pattern: '^(?P<label>49-[0-9]{{3}})$'
      label_only_heading_pattern: '^[A-Z ]+$'
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-07-03-scanned-rule",
    )

    assert ocr_calls == 1
    assert report.block_count == 2
    records = load_provisions(report.provisions_path)
    assert [record.citation_path for record in records] == [
        "us-test/regulation/scanned-rule",
        "us-test/regulation/scanned-rule/49-001",
        "us-test/regulation/scanned-rule/49-002",
    ]
    assert records[1].body == "OCR body text."
    assert records[2].body == "More OCR body text."


def test_extract_official_documents_can_force_ocr_when_pdf_has_page_number(
    tmp_path: Path, monkeypatch
) -> None:
    pdf_path = tmp_path / "scanned-table.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "1")
    document.save(pdf_path)
    document.close()
    ocr_calls = 0

    def fake_ocr_page(page, *, extraction):
        nonlocal ocr_calls
        del page
        ocr_calls += 1
        assert extraction["ocr"] is True
        assert extraction["force_ocr"] is True
        return "Table amount € 180,19"

    monkeypatch.setattr(documents_module, "_ocr_pdf_page_text", fake_ocr_page)
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: scanned-table
    jurisdiction: us-test
    document_class: guidance
    title: Scanned Table
    source_url: https://example.test/scanned-table.pdf
    citation_path: us-test/guidance/scanned-table
    source_format: pdf
    local_path: {json.dumps(str(pdf_path))}
    extraction:
      ocr: true
      force_ocr: true
      page_citation_prefix: schedule
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-07-04-scanned-table",
    )

    assert ocr_calls == 1
    assert report.block_count == 1
    records = load_provisions(report.provisions_path)
    assert records[1].citation_path == "us-test/guidance/scanned-table/schedule-1"
    assert records[1].body == "Table amount € 180,19"


def test_extract_labeled_pdf_sections_can_start_after_pattern(tmp_path: Path) -> None:
    pdf_path = tmp_path / "rules-with-toc.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "\n".join(
            [
                "TABLE OF CONTENTS",
                "49-001 First Section",
                "49-002 Second Section",
                "BEGIN RULE TEXT",
                "49-001 First Section. Actual text begins.",
                "Body text.",
                "49-002 Second Section. More text.",
            ]
        ),
    )
    document.save(pdf_path)
    document.close()
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: toc-rule
    jurisdiction: us-test
    document_class: regulation
    title: TOC Rule
    source_url: https://example.test/toc-rule.pdf
    citation_path: us-test/regulation/toc-rule
    source_format: pdf
    local_path: {json.dumps(str(pdf_path))}
    extraction:
      segmentation: labeled_sections
      start_after_pattern: '^BEGIN RULE TEXT$'
      section_heading_pattern: '^(?P<label>49-[0-9]{{3}})\\s+(?P<heading>.+)$'
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-07-03-toc-rule",
    )

    assert report.block_count == 2
    records = load_provisions(report.provisions_path)
    assert [record.heading for record in records[1:]] == [
        "49-001 First Section. Actual text begins.",
        "49-002 Second Section. More text.",
    ]
    assert records[1].body == "Body text."


def test_extract_official_documents_scrubs_public_mapbox_tokens(tmp_path: Path) -> None:
    html_path = tmp_path / "cms.html"
    public_token = "pk." + "abc_123" + "." + "DEF-456"
    html_path.write_text(
        f"""
        <html>
          <body>
            <main><h1>Medicare Eligibility</h1><p>People age 65 can get Medicare.</p></main>
            <script>{{"mapboxToken":"{public_token}"}}</script>
          </body>
        </html>
        """,
        encoding="utf-8",
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: cms-medicare
    jurisdiction: us
    document_class: guidance
    title: Medicare Eligibility
    source_url: https://www.cms.gov/example
    source_format: html
    local_path: {json.dumps(str(html_path))}
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-06-26-cms-medicare",
    )

    archived = report.source_paths[0].read_text(encoding="utf-8")
    assert public_token not in archived
    assert "[redacted-mapbox-public-token]" in archived
    assert load_provisions(report.provisions_path)[1].body == "People age 65 can get Medicare."


def test_extract_official_documents_uses_html_content_selector(tmp_path: Path) -> None:
    html_path = tmp_path / "wa-eaz.html"
    html_path.write_text(
        """
        <html>
          <head><title>Browser title</title></head>
          <body>
            <h1>Basic Food - Work Requirements</h1>
            <div class="site-shell">Navigation text should be ignored.</div>
            <div class="field-name-body">
              <p>Revised June 1, 2026</p>
              <h2>Purpose</h2>
              <p>Basic Food applicants must meet work registration rules.</p>
            </div>
          </body>
        </html>
        """,
        encoding="utf-8",
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: wa-eaz-work-requirements
    jurisdiction: us-wa
    document_class: manual
    title: Basic Food - Work Requirements
    source_url: https://www.dshs.wa.gov/esa/example
    source_format: html
    local_path: {json.dumps(str(html_path))}
    extraction:
      html_content_selector: .field-name-body
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-05-27-wa-eaz-manual",
    )

    assert report.block_count == 2
    records = load_provisions(report.provisions_path)
    bodies = "\n".join(record.body or "" for record in records)
    assert "Basic Food applicants must meet work registration rules" in bodies
    assert "Navigation text should be ignored" not in bodies


def test_extract_official_documents_segments_html_anchor_range(tmp_path: Path) -> None:
    html_path = tmp_path / "be-program-law.html"
    html_path.write_text(
        """
        <html>
          <body>
            <div id="list-title-3">
              <p>Earlier consolidated text.</p>
              <!-- field-start:statute -->
              <a name="Art.419">Art.</a> <a href="#Art.419bis">419</a>.
              <br>
              - droit d'accise : 245,4146 euros par 1 000 litres.
              <br>
              - droit d'accise special : 393,7887 euros par 1 000 litres.
              <!-- field-end:statute -->
              <a name="Art.420">Art.</a> <a href="#Art.420bis">420</a>.
              <br>
              Next article text should not be included.
            </div>
          </body>
        </html>
        """,
        encoding="utf-8",
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: be-program-law-article-419
    jurisdiction: be
    document_class: statute
    title: Belgian Program Law Article 419
    source_url: https://www.ejustice.just.fgov.be/eli/loi/2004/12/27/2004021170/justel
    source_format: html
    local_path: {json.dumps(str(html_path))}
    citation_path: be/statute/justel/excise/energy-products/2004021170
    extraction:
      html_content_selector: "#list-title-3"
      segmentation: anchor_range
      html_start_selector: 'a[name="Art.419"]'
      html_stop_selector: 'a[name="Art.420"]'
      section_label: Article 419
      section_heading: Article 419
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-06-30-be-excise",
    )

    assert report.block_count == 1
    records = load_provisions(report.provisions_path)
    section = next(record for record in records if record.kind == "section")
    assert section.citation_path == (
        "be/statute/justel/excise/energy-products/2004021170/Article 419"
    )
    assert "245,4146 euros" in (section.body or "")
    assert "393,7887 euros" in (section.body or "")
    assert "field-start:statute" not in (section.body or "")
    assert "field-end:statute" not in (section.body or "")
    assert "Next article text should not be included" not in (section.body or "")


def test_extract_official_documents_segments_multiple_html_anchor_ranges(
    tmp_path: Path,
) -> None:
    html_path = tmp_path / "be-alcohol-law.html"
    html_path.write_text(
        """
        <html>
          <body>
            <div id="list-title-3">
              <a name="Art.5">Art.</a> <a href="#Art.6">5</a>.
              Beer rate: 0,7933 EUR.
              <a name="Art.6">Art.</a> <a href="#Art.7">6</a>.
              Exemption text.
              <a name="Art.9">Art.</a> <a href="#Art.10">9</a>.
              Wine rate: 74,9086 EUR.
              <a name="Art.10">Art.</a> <a href="#Art.11">10</a>.
              Later article text.
            </div>
          </body>
        </html>
        """,
        encoding="utf-8",
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: be-alcohol-law
    jurisdiction: be
    document_class: statute
    title: Belgian Alcohol Excise Law
    source_url: https://www.ejustice.just.fgov.be/eli/loi/1998/01/07/1998003047/justel
    source_format: html
    local_path: {json.dumps(str(html_path))}
    citation_path: be/statute/justel/excise/alcohol/1998003047
    extraction:
      html_content_selector: "#list-title-3"
      segmentation: anchor_range
      anchor_ranges:
        - html_start_selector: 'a[name="Art.5"]'
          html_stop_selector: 'a[name="Art.6"]'
          section_label: Article 5
          section_heading: Article 5
        - html_start_selector: 'a[name="Art.9"]'
          html_stop_selector: 'a[name="Art.10"]'
          section_label: Article 9
          section_heading: Article 9
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-06-30-be-alcohol",
    )

    assert report.block_count == 2
    records = load_provisions(report.provisions_path)
    bodies_by_path = {
        record.citation_path: record.body for record in records if record.kind == "section"
    }
    article_5_body = bodies_by_path["be/statute/justel/excise/alcohol/1998003047/Article 5"]
    article_9_body = bodies_by_path["be/statute/justel/excise/alcohol/1998003047/Article 9"]
    assert "0,7933 EUR" in (article_5_body or "")
    assert "Exemption text" not in (article_5_body or "")
    assert "74,9086 EUR" in (article_9_body or "")
    assert "Later article text" not in (article_9_body or "")


def test_extract_official_documents_splits_labeled_html_sections(tmp_path: Path) -> None:
    html_path = tmp_path / "nm-nmac.html"
    html_path.write_text(
        """
        <html>
          <head><title>8.139.520 NMAC</title></head>
          <body>
            <div class="WordSection1">
              <p>TITLE 8 SOCIAL SERVICES</p>
              <p>8.139.520.1 ISSUING AGENCY: New Mexico Health Care Authority.</p>
              <p>[8.139.520.1 NMAC - Rp, 11/21/2023]</p>
              <p>8.139.520.2 SCOPE: General public.</p>
              <p>[8.139.520.2 NMAC - Rp, 11/21/2023]</p>
            </div>
          </body>
        </html>
        """,
        encoding="utf-8",
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: nm-nmac-8-139-520
    jurisdiction: us-nm
    document_class: regulation
    title: 8.139.520 NMAC Eligibility Policy - Income and Deductions
    source_url: https://www.srca.nm.gov/parts/title08/08.139.0520.html
    source_format: html
    local_path: {json.dumps(str(html_path))}
    citation_path: us-nm/regulation/nmac/8/139/520
    extraction:
      html_content_selector: .WordSection1
      segmentation: labeled_sections
      section_heading_pattern: '^(?P<label>8\\.139\\.520\\.\\d+)\\s+(?P<heading>[A-Z][A-Z ]+:)(?:\\s+(?P<body>.*))?$'
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-05-27-nm-snap-regulations",
    )

    assert report.block_count == 2
    records = load_provisions(report.provisions_path)
    assert [record.citation_path for record in records if record.kind == "section"] == [
        "us-nm/regulation/nmac/8/139/520/8.139.520.1",
        "us-nm/regulation/nmac/8/139/520/8.139.520.2",
    ]
    section_body = next(
        record.body for record in records if record.citation_path.endswith("8.139.520.1")
    )
    assert section_body == (
        "New Mexico Health Care Authority.\n\n[8.139.520.1 NMAC - Rp, 11/21/2023]"
    )


def test_extract_official_documents_formats_labeled_html_section_labels(
    tmp_path: Path,
) -> None:
    html_path = tmp_path / "nh-he-w-700.html"
    html_path.write_text(
        """
        <html>
          <body>
            <div class="WordSection1">
              <p>He-W 701 .02 Definitions F - O .</p>
              <p>(a) "Fraud" means an intentional program violation.</p>
              <p>He-W 702.03 Telephone Application .</p>
              <p>(a) Applicants may request assistance by telephone.</p>
              <p>APPENDIX A</p>
              <p>He-W 701.02 RSA 161:4-a, IV</p>
            </div>
          </body>
        </html>
        """,
        encoding="utf-8",
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: nh-he-w-700
    jurisdiction: us-nh
    document_class: regulation
    title: New Hampshire He-W 700 SNAP Rules
    source_url: https://gc.nh.gov/rules/state_agencies/he-w700.html
    source_format: html
    local_path: {json.dumps(str(html_path))}
    citation_path: us-nh/regulation/he-w-700
    extraction:
      html_content_selector: .WordSection1
      segmentation: labeled_sections
      section_heading_pattern: '^(?P<prefix>He-W)\\s+(?P<part>\\d+)\\s*\\.\\s*(?P<section>\\d+)\\s*(?P<heading>.+?)\\.?$'
      section_label_template: '{{prefix}} {{part}}.{{section}}'
      stop_text_pattern: '^APPENDIX A$'
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-05-27-nh-he-w-700-snap-rules",
    )

    records = load_provisions(report.provisions_path)
    assert [record.citation_path for record in records if record.kind == "section"] == [
        "us-nh/regulation/he-w-700/He-W 701.02",
        "us-nh/regulation/he-w-700/He-W 702.03",
    ]
    assert records[-1].body == "(a) Applicants may request assistance by telephone."


def test_extract_official_documents_keeps_content_inside_aspnet_form(
    tmp_path: Path,
) -> None:
    html_path = tmp_path / "il-dhs.html"
    html_path.write_text(
        """
        <html>
          <body>
            <form id="aspnetForm">
              <input type="hidden" name="__VIEWSTATE" value="ignored" />
              <div id="Main2">
                <h1>PM 13: SNAP Eligibility and Benefit Amount</h1>
                <p>SNAP eligibility rules are published in the manual body.</p>
              </div>
            </form>
          </body>
        </html>
        """,
        encoding="utf-8",
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: il-dhs-pm13
    jurisdiction: us-il
    document_class: manual
    title: "PM 13: SNAP Eligibility and Benefit Amount"
    source_url: https://www.dhs.state.il.us/page.aspx?item=16111
    source_format: html
    local_path: {json.dumps(str(html_path))}
    extraction:
      html_content_selector: "#Main2"
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-05-27-il-snap-manual",
    )

    assert report.block_count == 1
    records = load_provisions(report.provisions_path)
    bodies = "\n".join(record.body or "" for record in records)
    assert "SNAP eligibility rules are published in the manual body" in bodies
    assert "__VIEWSTATE" not in bodies


def test_extract_official_documents_preserves_content_after_unclosed_hidden_input(
    tmp_path: Path,
) -> None:
    html_path = tmp_path / "ssa-poms.html"
    html_path.write_text(
        """
        <!DOCTYPE html>
        <html>
          <body>
            <input id="Start" name="Start" type="hidden" value="">
            <div id="divBody">
              <h1>SI 01415.058</h1>
              <h2>District of Columbia</h2>
              <p>OS-A applies to recipients in Adult Foster Care Homes.</p>
            </div>
          </body>
        </html>
        """,
        encoding="utf-8",
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: ssa-poms-dc-ossp
    jurisdiction: us
    document_class: guidance
    title: "POMS SI 01415.058"
    source_url: https://secure.ssa.gov/apps10/poms.nsf/lnx/0501415058
    source_format: html
    local_path: {json.dumps(str(html_path))}
    extraction:
      html_content_selector: "#divBody"
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-07-03-dc-ossp-ssa-poms",
    )

    assert report.block_count == 1
    bodies = "\n".join(record.body or "" for record in load_provisions(report.provisions_path))
    assert "OS-A applies to recipients in Adult Foster Care Homes" in bodies
    assert "hidden" not in bodies


def test_extract_official_documents_from_json_html_field(tmp_path: Path) -> None:
    json_path = tmp_path / "ne-title-475-chapter-1.json"
    json_path.write_text(
        json.dumps(
            {
                "isSuccess": True,
                "output": {
                    "chapterHtml": (
                        "<p><strong>TITLE 475</strong></p>"
                        "<p><strong>CHAPTER 1 GENERAL PROVISIONS</strong></p>"
                        "<p>SNAP is a federal low income nutrition program.</p>"
                    )
                },
            }
        ),
        encoding="utf-8",
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: ne-title-475-chapter-1
    jurisdiction: us-ne
    document_class: regulation
    title: "Nebraska Title 475 NAC Chapter 1: General Provisions"
    source_url: https://rules.nebraska.gov/api/chapter/1741
    source_format: json
    local_path: {json.dumps(str(json_path))}
    extraction:
      json_html_field: output.chapterHtml
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-05-27-ne-snap-rules",
    )

    assert report.block_count == 1
    records = load_provisions(report.provisions_path)
    assert records[1].source_format == "json"
    assert "SNAP is a federal low income nutrition program" in (records[1].body or "")


def test_extract_official_documents_from_json_html_field_as_single_block(
    tmp_path: Path,
) -> None:
    json_path = tmp_path / "flanders-article.json"
    json_path.write_text(
        json.dumps(
            {
                "Tekst": (
                    "§ 1. Opening text before a paragraph.<br><p>§ 2. Later paragraph text.</p>"
                )
            }
        ),
        encoding="utf-8",
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: flanders-article
    jurisdiction: be-vlg
    document_class: statute
    title: "Flanders article"
    source_url: https://example.test/flanders-article
    source_format: json
    local_path: {json.dumps(str(json_path))}
    extraction:
      json_html_field: Tekst
      json_html_as_single_block: true
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-07-01-be-vlg-test",
    )

    assert report.block_count == 1
    records = load_provisions(report.provisions_path)
    assert "Opening text before a paragraph" in (records[1].body or "")
    assert "Later paragraph text" in (records[1].body or "")


def test_extract_official_documents_from_json_base64_html_field(
    tmp_path: Path,
) -> None:
    json_path = tmp_path / "fisconet-article.json"
    encoded_html = base64.b64encode(
        (
            "<html><body><h1>Article 48</h1>"
            "<p>Les droits sont perçus d'après le tarif.</p></body></html>"
        ).encode()
    ).decode("ascii")
    json_path.write_text(
        json.dumps({"data": {"content": {"content": encoded_html}}}),
        encoding="utf-8",
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: fisconet-article-48
    jurisdiction: be-bru
    document_class: statute
    title: "FisconetPlus article 48"
    source_url: https://example.test/fisconet-article-48
    source_format: json
    local_path: {json.dumps(str(json_path))}
    extraction:
      json_html_field: data.content.content
      json_html_base64: true
      json_html_as_single_block: true
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-07-01-be-fisconet-test",
    )

    assert report.block_count == 1
    records = load_provisions(report.provisions_path)
    assert "Les droits sont perçus d'après le tarif" in (records[1].body or "")


def test_extract_official_documents_from_json_records(tmp_path: Path) -> None:
    json_path = tmp_path / "ok-oac-340-50.json"
    json_path.write_text(
        json.dumps(
            [
                {
                    "id": 1,
                    "sectionNum": "340:50-1-1",
                    "description": "Purpose, legal base, and responsibilities",
                    "name": "Section",
                    "statusName": "Undefined",
                    "text": (
                        "<div>(a) <b>Purpose.</b> SNAP policy text.</div><div>(b) Other text.</div>"
                    ),
                },
                {
                    "id": 2,
                    "sectionNum": "340:50-1-2",
                    "description": "Legal basis",
                    "name": "Section",
                    "statusName": "Revoked",
                    "text": "<div>Old revoked text.</div>",
                },
                {
                    "id": 3,
                    "sectionNum": None,
                    "description": "General Provisions",
                    "name": "Subchapter",
                    "statusName": "Undefined",
                    "text": None,
                },
            ]
        ),
        encoding="utf-8",
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: ok-oac-340-50-snap
    jurisdiction: us-ok
    document_class: regulation
    title: "Oklahoma OAC 340:50 Supplemental Nutrition Assistance Program"
    source_url: https://rules.ok.gov/home
    source_format: json
    local_path: {json.dumps(str(json_path))}
    citation_path: us-ok/regulation/oac/340/50
    extraction:
      segmentation: records
      json_record_text_field: text
      json_record_text_is_html: true
      json_record_label_field: sectionNum
      json_record_heading_field: description
      json_record_kind_field: name
      json_record_status_field: statusName
      json_record_exclude_statuses:
        - Revoked
      json_record_metadata_fields:
        - id
        - sectionNum
        - statusName
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-05-27-ok-snap-rules",
    )

    assert report.block_count == 1
    records = load_provisions(report.provisions_path)
    assert [record.citation_path for record in records] == [
        "us-ok/regulation/oac/340/50",
        "us-ok/regulation/oac/340/50/340-50-1-1",
    ]
    assert records[1].heading == ("340:50-1-1 Purpose, legal base, and responsibilities")
    assert records[1].kind == "section"
    assert records[1].metadata is not None
    assert records[1].metadata["id"] == 1
    assert "SNAP policy text" in (records[1].body or "")
    assert "Old revoked text" not in (records[1].body or "")


def test_extract_official_documents_drops_configured_html_selectors(
    tmp_path: Path,
) -> None:
    html_path = tmp_path / "ks-keesm.html"
    html_path.write_text(
        """
        <html>
          <head><title>4200 Assistance Planning</title></head>
          <body>
            <h1>Kansas Economic and Employment Services Manual</h1>
            <h2>4000 Assistance Planning</h2>
            <h6>04-26</h6>
            <p>Food assistance household composition rules apply.</p>
          </body>
        </html>
        """,
        encoding="utf-8",
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: ks-keesm-4200
    jurisdiction: us-ks
    document_class: manual
    title: "Kansas KEESM: 4200 Assistance Planning"
    source_url: https://content.dcf.ks.gov/EES/KEESM/Current/keesm4200.htm
    source_format: html
    local_path: {json.dumps(str(html_path))}
    extraction:
      html_drop_selectors:
        - h1
        - h6
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-05-27-ks-keesm",
    )

    assert report.block_count == 1
    records = load_provisions(report.provisions_path)
    block = records[1]
    assert block.heading == "4000 Assistance Planning"
    assert "Food assistance household composition rules apply" in (block.body or "")
    assert "04-26" not in (block.body or "")


def test_extract_official_documents_reads_docx_sections(tmp_path: Path) -> None:
    docx_path = tmp_path / "chapter-365.docx"
    document_xml = """\
<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p>
      <w:pPr><w:pStyle w:val="Heading1"/></w:pPr>
      <w:r><w:t>365.100 Special Situation Households</w:t></w:r>
    </w:p>
    <w:p><w:r><w:t>SNAP households may include special living arrangements.</w:t></w:r></w:p>
    <w:tbl>
      <w:tr>
        <w:tc><w:p><w:r><w:t>Household</w:t></w:r></w:p></w:tc>
        <w:tc><w:p><w:r><w:t>Treatment</w:t></w:r></w:p></w:tc>
      </w:tr>
    </w:tbl>
    <w:p>
      <w:pPr><w:pStyle w:val="Heading2"/></w:pPr>
      <w:r><w:t>365.110 Residents of Institutions</w:t></w:r>
    </w:p>
    <w:p><w:r><w:t>Institution residents are subject to program rules.</w:t></w:r></w:p>
  </w:body>
</w:document>
"""
    with zipfile.ZipFile(docx_path, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: ma-dta-chapter-365
    jurisdiction: us-ma
    document_class: regulation
    title: "Massachusetts DTA Chapter 365"
    source_url: https://www.mass.gov/doc/chapter-365-special-situation-households/download
    source_format: docx
    local_path: {json.dumps(str(docx_path))}
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-05-27-ma-dta-regulations",
    )

    assert report.block_count == 2
    records = load_provisions(report.provisions_path)
    first_block = records[1]
    assert first_block.heading == "365.100 Special Situation Households"
    assert "SNAP households may include special living arrangements" in (first_block.body or "")
    assert "Household | Treatment" in (first_block.body or "")
    second_block = records[2]
    assert second_block.heading == "365.110 Residents of Institutions"


def test_extract_official_documents_reads_legacy_word_doc(tmp_path: Path, monkeypatch) -> None:
    doc_path = tmp_path / "5030_10.doc"
    doc_path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1legacy-word-content")
    monkeypatch.setattr(
        documents_module,
        "_legacy_word_document_text",
        lambda content: (
            "5030.10 Earned Income Disregard\n\n"
            "$65.00 per month plus 1/2 of the remaining income is disregarded."
        ),
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: ct-upm-5030-10
    jurisdiction: us-ct
    document_class: policy
    title: "UPM 5030.10 - Earned Income Disregards"
    source_url: https://portal.ct.gov/dss/-/media/departments-and-agencies/dss/upms/upm5---treatment-of-income-income-eligibility/5030_10.doc
    source_format: doc
    local_path: {json.dumps(str(doc_path))}
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-07-02-ct-upm-ssp",
    )

    assert report.block_count == 1
    records = load_provisions(report.provisions_path)
    assert records[1].source_format == "doc"
    assert records[1].heading == "UPM 5030.10 - Earned Income Disregards"
    assert "$65.00 per month plus 1/2" in (records[1].body or "")
    assert len(report.source_paths) == 1
    assert report.source_paths[0].suffix == ".doc"
    assert report.source_paths[0].read_bytes().startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")


def test_extract_official_documents_segments_labeled_docx_sections(
    tmp_path: Path,
) -> None:
    docx_path = tmp_path / "chapter-704.docx"
    document_xml = """\
<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:r><w:t>TABLE OF CONTENTS</w:t></w:r></w:p>
    <w:p><w:r><w:t>704.000 Overview of Financial Eligibility</w:t></w:r></w:p>
    <w:p><w:r><w:t>704.500 Calculation of Grant Amount</w:t></w:r></w:p>
    <w:p><w:r><w:t>106 CMR: Department of Transitional Assistance | Page | 704.000</w:t></w:r></w:p>
    <w:p><w:r><w:t>704.000: Overview of Financial Eligibility</w:t></w:r></w:p>
    <w:p><w:r><w:t>Applicants must meet financial eligibility requirements.</w:t></w:r></w:p>
    <w:p><w:r><w:t>106 CMR: Department of Transitional Assistance</w:t></w:r></w:p>
    <w:p><w:r><w:t>704.500: Calculation of Grant Amount</w:t></w:r></w:p>
    <w:p><w:r><w:t>Step 1: Identify earned income.</w:t></w:r></w:p>
    <w:p><w:r><w:t>Step 2: Subtract disregarded income.</w:t></w:r></w:p>
  </w:body>
</w:document>
"""
    with zipfile.ZipFile(docx_path, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: ma-dta-chapter-704
    jurisdiction: us-ma
    document_class: regulation
    title: "Massachusetts DTA Chapter 704"
    source_url: https://www.mass.gov/doc/chapter-704-financial-eligibility-0/download
    source_format: docx
    local_path: {json.dumps(str(docx_path))}
    citation_path: us-ma/regulation/dta/106-cmr/704
    extraction:
      segmentation: labeled_sections
      start_after_pattern: 'Page\\s+\\|\\s+704\\.000$'
      section_heading_pattern: '^(?P<label>704\\.\\d{{3}}):?\\s+(?P<heading>.+)$'
      drop_line_patterns:
        - '^106 CMR:'
        - '^TABLE OF CONTENTS$'
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-06-27-ma-tafdc-regulations",
    )

    assert report.block_count == 2
    records = load_provisions(report.provisions_path)
    assert [record.citation_path for record in records if record.kind == "section"] == [
        "us-ma/regulation/dta/106-cmr/704/704.000",
        "us-ma/regulation/dta/106-cmr/704/704.500",
    ]
    assert records[1].body == "Applicants must meet financial eligibility requirements."
    assert records[2].body == (
        "Step 1: Identify earned income.\n\nStep 2: Subtract disregarded income."
    )


def test_extract_official_documents_reads_webworks_policy_divs(tmp_path: Path) -> None:
    html_path = tmp_path / "az-faa5.html"
    html_path.write_text(
        """
        <html>
          <head><title>A CA Benefit Determination</title></head>
          <body>
            <div id="ww_content_container">
              <div id="page_content_container">
                <div id="page_content">
                  <div class="Heading_Subject">CA Benefit Determination</div>
                  <div class="WebWorks_MiniTOC">
                    <a class="WebWorks_MiniTOC_Link" href="#policy">Policy</a>
                  </div>
                  <div class="Body_Text_Public">
                    This section identifies how benefit amounts are determined.
                  </div>
                  <div class="Heading_Section_Public">Policy</div>
                  <div class="Body_Text_Public">
                    Countable income is compared to the payment standard.
                  </div>
                  <div class="List_Bullet_Public">
                    <span class="WebWorks_Number">-</span> Net income is subtracted.
                  </div>
                  <div class="Heading_Section_Public">Examples</div>
                  <div class="ww_skin_page_overflow">
                    <table>
                      <tr>
                        <td><div class="Body_Text_Bullet_Public">Individual</div></td>
                        <td><div class="Body_Text_Bullet_Public">Family Member</div></td>
                      </tr>
                    </table>
                  </div>
                </div>
              </div>
            </div>
          </body>
        </html>
        """,
        encoding="utf-8",
    )
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: az-faa5-ca-benefit
    jurisdiction: us-az
    document_class: manual
    title: Arizona DES FAA5 CA Benefit Determination
    source_url: https://dbmefaapolicy.azdes.gov/FAA5/CA_Benefit_Determination.html
    source_format: html
    local_path: {json.dumps(str(html_path))}
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2025-10-30-az-des-faa5-manual",
    )

    assert report.block_count == 3
    records = load_provisions(report.provisions_path)
    policy_block = next(record for record in records if record.heading == "Policy")
    policy_body = policy_block.body or ""
    assert "Countable income is compared to the payment standard" in policy_body
    assert "Net income is subtracted" in policy_body
    assert "WebWorks_MiniTOC" not in policy_body
    examples_block = next(record for record in records if record.heading == "Examples")
    assert "Individual" in (examples_block.body or "")


def test_extract_official_documents_splits_numbered_pdf_sections(tmp_path):
    pdf_path = tmp_path / "rules.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "\n".join(
            [
                "IDAHO ADMINISTRATIVE CODE",
                "Section 000",
                "Page 1",
                "000.",
                "LEGAL AUTHORITY.",
                "Legal authority text.",
                "001.",
                "SCOPE.",
                "01.",
                "Scope body.",
                "002. -- 009.",
                "(RESERVED)",
            ]
        ),
    )
    document.save(pdf_path)
    document.close()
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: idaho-rule
    jurisdiction: us-id
    document_class: regulation
    title: IDAPA 35.01.01 - Income Tax Administrative Rules
    source_url: https://adminrules.idaho.gov/rules/current/35/350101.pdf
    citation_path: us-id/regulation/idapa/35/01/01
    source_format: pdf
    local_path: {json.dumps(str(pdf_path))}
    extraction:
      segmentation: numbered_sections
      drop_lines:
        - IDAHO ADMINISTRATIVE CODE
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-05-12-idapa-35-01-01",
    )

    assert report.block_count == 3
    records = load_provisions(report.provisions_path)
    assert [record.citation_path for record in records] == [
        "us-id/regulation/idapa/35/01/01",
        "us-id/regulation/idapa/35/01/01/000",
        "us-id/regulation/idapa/35/01/01/001",
        "us-id/regulation/idapa/35/01/01/002-009",
    ]
    first_section = records[1]
    assert first_section.kind == "section"
    assert first_section.heading == "000. LEGAL AUTHORITY."
    assert first_section.body == "Legal authority text."
    assert first_section.metadata is not None
    assert first_section.metadata["page_start"] == 1
    assert "section_end_label" not in first_section.metadata
    reserved_section = records[3]
    assert reserved_section.metadata is not None
    assert reserved_section.metadata["section_end_label"] == "009"


def test_extract_official_documents_splits_labeled_pdf_sections(tmp_path):
    pdf_path = tmp_path / "capi.pdf"
    document = fitz.open()
    first_page = document.new_page()
    first_page.insert_text(
        (72, 72),
        "\n".join(
            [
                "ELIGIBILITY AND ASSISTANCE STANDARDS",
                "49-001",
                "PROGRAM DEFINITION",
                ".1 Program definition text.",
                "CALIFORNIA-DSS-MANUAL-EAS",
                "Page 660.2",
            ]
        ),
    )
    second_page = document.new_page()
    second_page.insert_text(
        (72, 72),
        "\n".join(
            [
                "ELIGIBILITY AND ASSISTANCE STANDARDS",
                "49-001",
                "PROGRAM DEFINITION",
                ".2 Continued definition text.",
                "49-010",
                "ELIGIBILITY FOR CASH ASSISTANCE PROGRAM",
                "FOR IMMIGRANTS",
                "HANDBOOK BEGINS HERE",
                ".1 Eligibility text.",
            ]
        ),
    )
    document.save(pdf_path)
    document.close()
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: ca-capi
    jurisdiction: us-ca
    document_class: regulation
    title: California CAPI Regulations
    source_url: https://cdss.ca.gov/Portals/9/CAPI/CAPI_Regulations-Accessible.pdf
    citation_path: us-ca/regulation/cdss/eas/49
    source_format: pdf
    local_path: {json.dumps(str(pdf_path))}
    extraction:
      segmentation: labeled_sections
      section_label_pattern: "^(?P<label>49-[0-9]{{3}})$"
      drop_lines:
        - ELIGIBILITY AND ASSISTANCE STANDARDS
        - CALIFORNIA-DSS-MANUAL-EAS
      drop_line_patterns:
        - "^Page \\\\d+(?:\\\\.\\\\d+)?$"
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-05-12-capi",
    )

    assert report.block_count == 2
    records = load_provisions(report.provisions_path)
    assert [record.citation_path for record in records] == [
        "us-ca/regulation/cdss/eas/49",
        "us-ca/regulation/cdss/eas/49/49-001",
        "us-ca/regulation/cdss/eas/49/49-010",
    ]
    definition = records[1]
    assert definition.body == ".1 Program definition text. .2 Continued definition text."
    eligibility = records[2]
    assert eligibility.heading == "49-010 ELIGIBILITY FOR CASH ASSISTANCE PROGRAM FOR IMMIGRANTS"
    assert eligibility.body == "HANDBOOK BEGINS HERE .1 Eligibility text."
    assert eligibility.metadata is not None
    assert eligibility.metadata["page_start"] == 2


def test_extract_official_documents_replaces_labeled_pdf_section_labels(tmp_path):
    pdf_path = tmp_path / "roman.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "\n".join(
            [
                "I. FIRST POLICY HEADING",
                "First body.",
                "II. SECOND POLICY HEADING",
                "Second body.",
            ]
        ),
    )
    document.save(pdf_path)
    document.close()
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: roman-policy
    jurisdiction: us-nc
    document_class: manual
    title: Roman Policy
    source_url: https://policies.ncdhhs.gov/example.pdf
    citation_path: us-nc/manual/dhhs/work-first/roman-policy
    source_format: pdf
    local_path: {json.dumps(str(pdf_path))}
    extraction:
      segmentation: labeled_sections
      section_heading_pattern: "^(?P<label>I|II)\\\\.\\\\s+(?P<heading>.+)$"
      section_label_replacements:
        I: section-i
        II: section-ii
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-07-03-roman",
    )

    assert report.block_count == 2
    records = load_provisions(report.provisions_path)
    assert [record.citation_path for record in records] == [
        "us-nc/manual/dhhs/work-first/roman-policy",
        "us-nc/manual/dhhs/work-first/roman-policy/section-i",
        "us-nc/manual/dhhs/work-first/roman-policy/section-ii",
    ]
    assert records[1].heading == "section-i FIRST POLICY HEADING"
    assert records[2].body == "Second body."


def test_extract_official_documents_respects_labeled_pdf_end_page(tmp_path):
    pdf_path = tmp_path / "rules.pdf"
    document = fitz.open()
    first_page = document.new_page()
    first_page.insert_text(
        (72, 72),
        "\n".join(
            [
                "49-001",
                "FIRST SECTION",
                ".1 First section text.",
            ]
        ),
    )
    second_page = document.new_page()
    second_page.insert_text(
        (72, 72),
        "\n".join(
            [
                "49-002",
                "SECOND SECTION",
                ".1 Second section text.",
            ]
        ),
    )
    document.save(pdf_path)
    document.close()
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: bounded-rule
    jurisdiction: us-az
    document_class: regulation
    title: Bounded Rules
    source_url: https://apps.azsos.gov/public_services/Title_06/6-05.pdf
    citation_path: us-az/regulation/aac/title-6/chapter-5/article-49
    source_format: pdf
    local_path: {json.dumps(str(pdf_path))}
    extraction:
      segmentation: labeled_sections
      section_label_pattern: "^(?P<label>49-[0-9]{{3}})$"
      label_only_heading_pattern: "^[A-Z ]+$"
      start_page: 1
      end_page: 1
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-05-12-bounded",
    )

    assert report.block_count == 1
    records = load_provisions(report.provisions_path)
    assert [record.citation_path for record in records] == [
        "us-az/regulation/aac/title-6/chapter-5/article-49",
        "us-az/regulation/aac/title-6/chapter-5/article-49/49-001",
    ]
    assert records[1].body == ".1 First section text."


def test_extract_labeled_pdf_sections_formats_section_labels(
    tmp_path: Path,
) -> None:
    pdf_path = tmp_path / "source-book.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "\n".join(
            [
                "SECTION 1 \u2013 FORWARD/NOTICE",
                "SECTION 1: Forward/Notice",
                "Forward text.",
                "SECTION 2: Eligibility",
                "Eligibility text.",
            ]
        ),
    )
    document.save(pdf_path)
    document.close()
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: ny-source-book
    jurisdiction: us-ny
    document_class: manual
    title: New York SNAP Source Book
    source_url: https://otda.ny.gov/programs/snap/SNAPSB.pdf
    citation_path: us-ny/manual/otda/snap-source-book
    source_format: pdf
    local_path: {json.dumps(str(pdf_path))}
    extraction:
      segmentation: labeled_sections
      section_heading_pattern: "^(?P<label>SECTION\\\\s+(?P<number>\\\\d+))(?::|\\\\s+[\\\\u2013-]\\\\s+)\\\\s*(?P<heading>.+)$"
      section_label_template: "section-{{number}}"
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-05-27-ny-source-book",
    )

    assert report.coverage.complete
    records = load_provisions(report.provisions_path)
    assert [record.citation_path for record in records] == [
        "us-ny/manual/otda/snap-source-book",
        "us-ny/manual/otda/snap-source-book/section-1",
        "us-ny/manual/otda/snap-source-book/section-2",
    ]
    assert records[1].metadata is not None
    assert records[1].metadata["section_label"] == "section-1"


def test_extract_labeled_pdf_sections_supports_label_heading_next_line(
    tmp_path: Path,
) -> None:
    pdf_path = tmp_path / "delaware.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "\n".join(
            [
                "9001",
                "Legal Base",
                "Legal base text.",
                "9002 Penalties",
                "Penalty text.",
                "9064.7 Households Entitled to a Deduction in DSSM",
                "9060",
                "[273.10(d)(7)]",
                "Continuation text.",
                "9065",
                "Calculating Net Income",
                "Net income text.",
            ]
        ),
    )
    document.save(pdf_path)
    document.close()
    manifest_path = tmp_path / "documents.yaml"
    manifest_path.write_text(
        f"""
documents:
  - source_id: de-snap-rules
    jurisdiction: us-de
    document_class: regulation
    title: Delaware SNAP Rules
    source_url: https://regulations.delaware.gov/AdminCode/title16/9000
    citation_path: us-de/regulation/title-16/9000-food-stamp-program
    source_format: pdf
    local_path: {json.dumps(str(pdf_path))}
    extraction:
      segmentation: labeled_sections
      section_heading_pattern: '^(?P<label>9\\d{{3}}(?:\\.\\d+)*)\\s+(?P<heading>[A-Z][A-Za-z0-9 ().,''/-]+)$'
      section_label_pattern: '^(?P<label>9\\d{{3}}(?:\\.\\d+)*)$'
      label_only_heading_pattern: '^[A-Z][A-Za-z0-9 ().,''/-]+$'
      label_only_requires_heading: true
"""
    )
    store = CorpusArtifactStore(tmp_path / "corpus")

    report = extract_official_documents(
        store,
        manifest_path=manifest_path,
        version="2026-05-27-de-snap-rules",
    )

    assert report.coverage.complete
    records = load_provisions(report.provisions_path)
    assert [record.citation_path for record in records] == [
        "us-de/regulation/title-16/9000-food-stamp-program",
        "us-de/regulation/title-16/9000-food-stamp-program/9001",
        "us-de/regulation/title-16/9000-food-stamp-program/9002",
        "us-de/regulation/title-16/9000-food-stamp-program/9064.7",
        "us-de/regulation/title-16/9000-food-stamp-program/9065",
    ]
    assert records[1].heading == "9001 Legal Base"
    assert "9060 [273.10(d)(7)] Continuation text." in (records[3].body or "")
